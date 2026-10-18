"""Content extraction from course documents.

Every extractor returns a list of :class:`ExtractedPage` carrying a **1-based**
``page_number`` so that any derived chunk can be cited back to an exact PDF
page or PPTX slide.

Supported formats
-----------------
========  ==========================================================
PDF       :mod:`pdfplumber` (layout-aware), falling back to ``pypdf``
DOCX      :mod:`python-docx` — real pages tracked via ``<w:br type=page>``
PPTX      :mod:`python-pptx` — one entry per slide
TXT / MD  split on form feeds (``\\f``) when present, else one page
========  ==========================================================
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".pptx", ".txt", ".md"})


class ExtractionError(Exception):
    """Raised when a document cannot be parsed into text."""


@dataclass(frozen=True, slots=True)
class ExtractedPage:
    """One page (or slide) of extracted text."""

    page_number: int  # 1-based
    text: str

    @property
    def word_count(self) -> int:
        return len(self.text.split())

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


# ---------------------------------------------------------------------------
# Text cleaning
# ---------------------------------------------------------------------------

_NEWLINES = re.compile(r"\r\n?|\x0b")
_HORIZ_WS = re.compile(r"[^\S\n]+")
_BLANK_LINES = re.compile(r"\n{3,}")


def clean_text(text: str) -> str:
    """Normalise whitespace without touching the wording.

    * ``\\r\\n`` / ``\\r`` / ``\\v``  -> ``\\n``
    * runs of spaces/tabs           -> a single space
    * 3+ consecutive newlines       -> exactly one blank line
    * trailing whitespace           -> stripped per line

    Arabic letters, diacritics and the ``\\u200f``/``\\u200e`` marks are left
    untouched so that bidirectional text keeps rendering correctly.
    """

    if not text:
        return ""

    text = _NEWLINES.sub("\n", text)
    text = _HORIZ_WS.sub(" ", text)
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = _BLANK_LINES.sub("\n\n", text)

    return text.strip()


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------


def extract_pages(path: Path | str) -> list[ExtractedPage]:
    """Extract ``path`` into an ordered list of non-empty pages."""

    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            f"Unsupported file type {suffix!r}. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    if not file_path.is_file():
        raise ExtractionError(f"File not found: {file_path}")

    extractors = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".pptx": _extract_pptx,
        ".txt": _extract_plaintext,
        ".md": _extract_plaintext,
    }

    try:
        pages = extractors[suffix](file_path)
    except ExtractionError:
        raise
    except Exception as exc:  # pragma: no cover - defensive
        raise ExtractionError(f"Failed to extract {file_path.name}: {exc}") from exc

    # Never return empty pages -- they would pollute the chunk index.
    return [page for page in pages if not page.is_empty]


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _extract_pdf(path: Path) -> list[ExtractedPage]:
    """Page-by-page PDF extraction via pdfplumber, with a pypdf fallback."""

    try:
        return _extract_pdf_plumber(path)
    except Exception:
        # Scanned/image-heavy PDFs sometimes break pdfplumber's layout engine.
        return _extract_pdf_pypdf(path)


def _extract_pdf_plumber(path: Path) -> list[ExtractedPage]:
    import pdfplumber

    pages: list[ExtractedPage] = []

    with pdfplumber.open(path) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            # Default (non-layout) mode gives the cleanest reading order for
            # single-column lecture slides; `clean_text` handles leftover gaps.
            raw = page.extract_text() or ""
            pages.append(ExtractedPage(page_number=index, text=clean_text(raw)))

    return pages


def _extract_pdf_pypdf(path: Path) -> list[ExtractedPage]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[ExtractedPage] = []

    for index, page in enumerate(reader.pages, start=1):
        raw = page.extract_text() or ""
        pages.append(ExtractedPage(page_number=index, text=clean_text(raw)))

    return pages


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _extract_docx(path: Path) -> list[ExtractedPage]:
    """Extract a DOCX, honouring explicit page breaks for page numbering.

    Word documents have no fixed pagination, so the only reliable way to cite a
    page is to count ``<w:br w:type="page"/>`` elements. Without any break the
    whole document is page 1.
    """

    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))

    pages: list[list[str]] = [[]]
    current = pages[0]

    for block in _iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            for run in block.runs:
                if _run_has_page_break(run):
                    current = []
                    pages.append(current)
            text = block.text.strip()
            if text:
                current.append(text)

        elif isinstance(block, Table):
            current.append(_table_to_text(block))

    result: list[ExtractedPage] = []
    for index, lines in enumerate(pages, start=1):
        text = clean_text("\n".join(lines))
        if text:
            result.append(ExtractedPage(page_number=index, text=text))

    return result


def _iter_docx_blocks(document):
    """Yield paragraphs and tables in true document order."""

    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _run_has_page_break(run) -> bool:
    """Detect ``<w:br w:type="page"/>`` inside a run."""

    from docx.oxml.ns import qn

    for br in run._element.findall(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def _table_to_text(table) -> str:
    """Render a DOCX/PPTX table as plain rows so its text is searchable."""

    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append(" | ".join(cells))
    return "\n".join(rows)


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def _extract_pptx(path: Path) -> list[ExtractedPage]:
    """Extract one :class:`ExtractedPage` per slide, in slide order."""

    from pptx import Presentation

    presentation = Presentation(str(path))
    pages: list[ExtractedPage] = []

    for index, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []

        for shape in _iter_slide_shapes(slide):
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    blocks.append(text)

            if getattr(shape, "has_table", False):
                table_text = _table_to_text(shape.table)
                if table_text.strip():
                    blocks.append(table_text)

        text = clean_text("\n\n".join(blocks))
        if text:
            pages.append(ExtractedPage(page_number=index, text=text))

    return pages


def _iter_slide_shapes(slide):
    """Yield shapes depth-first, so grouped shapes are not skipped."""

    for shape in slide.shapes:
        yield shape
        if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
            yield from _iter_slide_shapes(shape)


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------


def _extract_plaintext(path: Path) -> list[ExtractedPage]:
    """Split a text file on form feeds; otherwise treat it as a single page."""

    raw = path.read_text(encoding="utf-8", errors="replace")
    raw = _NEWLINES.sub("\n", raw)

    segments = raw.split("\f")

    pages: list[ExtractedPage] = []
    for index, segment in enumerate(segments, start=1):
        text = clean_text(segment)
        if text:
            pages.append(ExtractedPage(page_number=index, text=text))

    return pages
