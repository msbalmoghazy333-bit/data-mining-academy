"""Builders for sample documents used by the extraction tests.

``build_pdf`` writes a minimal but spec-valid PDF from scratch (correct byte
offsets, correct xref table) so the tests need **no** extra dependency such as
reportlab, and so page content is known exactly.
"""

from __future__ import annotations

import io
from pathlib import Path


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _content_stream(text: str) -> bytes:
    """Render ``text`` as a PDF content stream, one text line per input line."""

    lines = text.split("\n")
    parts = ["BT", "/F1 12 Tf", "72 720 Td", "16 TL"]

    for index, line in enumerate(lines):
        parts.append(f"({_pdf_escape(line)}) Tj")
        if index < len(lines) - 1:
            parts.append("T*")

    parts.append("ET")
    return "\n".join(parts).encode("ascii")


def build_pdf(pages: list[str]) -> bytes:
    """Build a valid PDF with exactly one page per entry in ``pages``."""

    if not pages:
        raise ValueError("build_pdf needs at least one page")

    count = len(pages)
    font_obj = 3
    first_page_obj = 4

    page_objs = [first_page_obj + 2 * i for i in range(count)]
    content_objs = [first_page_obj + 2 * i + 1 for i in range(count)]

    objects: dict[int, bytes] = {}

    kids = " ".join(f"{num} 0 R" for num in page_objs)
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode("ascii")
    objects[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    for index, text in enumerate(pages):
        stream = _content_stream(text)

        objects[page_objs[index]] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_obj} 0 R >> >> "
            f"/Contents {content_objs[index]} 0 R >>"
        ).encode("ascii")

        objects[content_objs[index]] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode("ascii")
            + stream
            + b"\nendstream"
        )

    # --- serialise with exact byte offsets ---------------------------------
    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    max_obj = 3 + 2 * count

    for number in range(1, max_obj + 1):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n".encode("ascii")
        out += objects[number]
        out += b"\nendobj\n"

    xref_offset = len(out)

    out += f"xref\n0 {max_obj + 1}\n".encode("ascii")
    # Every xref entry must be exactly 20 bytes.
    out += b"0000000000 65535 f \n"
    for number in range(1, max_obj + 1):
        out += f"{offsets[number]:010d} 00000 n \n".encode("ascii")

    out += (
        f"trailer\n<< /Size {max_obj + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    ).encode("ascii")

    return bytes(out)


def write_pdf(path: Path, pages: list[str]) -> Path:
    path.write_bytes(build_pdf(pages))
    return path


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def write_docx(path: Path, paragraphs: list[str], *, page_break_before: int | None = None) -> Path:
    """Write a DOCX.

    ``page_break_before`` inserts an explicit page break before that
    (0-based) paragraph, which the extractor must honour as a new page.
    """

    import docx

    document = docx.Document()

    for index, text in enumerate(paragraphs):
        paragraph = document.add_paragraph()
        if page_break_before is not None and index == page_break_before:
            paragraph.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
        paragraph.add_run(text)

    document.save(str(path))
    return path


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def write_pptx(path: Path, slides: list[list[str]]) -> Path:
    """Write a PPTX with one text-box line-group per slide."""

    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()

    for lines in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(6))
        frame = box.text_frame
        frame.text = lines[0]
        for extra in lines[1:]:
            frame.add_paragraph().text = extra

    buffer = io.BytesIO()
    presentation.save(buffer)
    path.write_bytes(buffer.getvalue())
    return path


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------


def write_txt(path: Path, pages: list[str]) -> Path:
    """Write a text file with pages separated by form feeds."""

    path.write_text("\f".join(pages), encoding="utf-8")
    return path
